"""
Files API — استقبال مرفقات لوحة التحكم وحفظها.
يحفظ في ~/.cobweaverclaw/uploads/ باسم موقوت، ويستخرج نص الملفات
النصية لحقنها في سياق المحادثة. حد الحجم 200MB للملف الواحد.
"""
from __future__ import annotations

import base64
import datetime
import logging
import re
from pathlib import Path
from typing import Any, Dict, List, Tuple

logger = logging.getLogger(__name__)

UPLOADS_DIR = Path.home() / ".cobweaverclaw" / "uploads"
MAX_FILE_BYTES = 200 * 1024 * 1024  # 200MB

# امتدادات تُقرأ نصاً وتُحقن في سياق النموذج
_TEXT_EXTS = {
    "txt", "md", "markdown", "csv", "json", "yaml", "yml", "xml", "html",
    "py", "js", "ts", "jsx", "tsx", "cpp", "c", "h", "hpp", "java", "go",
    "rs", "rb", "php", "sh", "bash", "sql", "toml", "ini", "cfg", "log",
}
_TEXT_INJECT_LIMIT = 12000  # حرف لكل ملف داخل السياق


def _safe_name(name: str) -> str:
    """اسم ملف آمن (يمنع اجتياز المسارات) مع الإبقاء على العربية."""
    name = (name or "file").replace("\\", "/").split("/")[-1]
    name = re.sub(r'[<>:"|?*\x00-\x1f]', "_", name).strip(". ") or "file"
    return name[:120]


def _ext(name: str) -> str:
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


def save_attachments(attachments: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """يفكّ base64 ويحفظ كل مرفق باسم موقوت. يعيد قائمة نتائج لكل ملف.

    كل نتيجة: {name, path, size, type, saved, error?, text?}
    text يُملأ للملفات النصية (لحقنها في سياق النموذج).
    """
    UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
    out: List[Dict[str, Any]] = []
    for att in attachments or []:
        if not isinstance(att, dict):
            continue
        name = _safe_name(str(att.get("name", "file")))
        rec: Dict[str, Any] = {"name": name, "type": str(att.get("type", "")),
                               "saved": False, "path": None,
                               "size": int(att.get("size") or 0)}
        try:
            raw_b64 = str(att.get("data", ""))
            # اقبل صيغة data URL أيضاً
            if "," in raw_b64 and raw_b64.lstrip().startswith("data:"):
                raw_b64 = raw_b64.split(",", 1)[1]
            blob = base64.b64decode(raw_b64, validate=False)
            rec["size"] = len(blob)
            if len(blob) > MAX_FILE_BYTES:
                rec["error"] = "الملف يتجاوز الحد الأقصى 200MB"
                out.append(rec)
                continue
            ts = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
            dest = UPLOADS_DIR / f"{ts}_{name}"
            i = 1
            while dest.exists():
                dest = UPLOADS_DIR / f"{ts}_{i}_{name}"
                i += 1
            dest.write_bytes(blob)
            rec["saved"] = True
            rec["path"] = str(dest)
            rec["saved_name"] = dest.name
            # استخراج المحتوى حسب النوع (نص/PDF/DOCX/جداول)
            ext_res = extract_content(dest, rec["type"])
            if ext_res.get("image"):
                rec["image"] = True
                rec["data"] = raw_b64   # يبقى للنموذج (vision) حتى في مسار الحفظ المتأخر
            else:
                content = ext_res.get("content") or ""
                rec["content"] = content
                rec["text"] = content[:_TEXT_INJECT_LIMIT]  # للحقن في سياق النموذج
                rec["pages"] = ext_res.get("pages", 0)
                if ext_res.get("error"):
                    rec["extract_error"] = ext_res["error"]
            logger.info("Attachment saved: %s (%d bytes)", dest.name, len(blob))
        except Exception as e:
            rec["error"] = f"فشل الحفظ: {str(e)[:80]}"
        out.append(rec)
    return out


def extract_content(file_path, mime_type: str = "") -> Dict[str, Any]:
    """يستخرج محتوى الملف نصاً حسب نوعه.

    يعيد {content, pages, error?} — الصور تعيد {image: True} بلا نص.
    مكتبات PDF/DOCX/XLSX تُستورد كسولاً؛ غيابها يعيد خطأً واضحاً بأمر التثبيت.
    """
    p = Path(file_path)
    name = p.name.lower()
    ext = _ext(name)
    mime = (mime_type or "").lower()
    out: Dict[str, Any] = {"content": "", "pages": 0}

    try:
        # صور — لا استخراج؛ تُرسل base64 للنموذج مباشرة
        if mime.startswith("image/") or ext in ("png", "jpg", "jpeg", "gif", "webp", "bmp", "svg"):
            out["image"] = True
            return out

        # جداول أولاً (قبل فرع النص العام — csv نصي لكنه يُقرأ كجدول)
        if ext in ("xlsx", "xls", "csv") or mime == "text/csv" \
           or mime.endswith("spreadsheetml.sheet"):
            pass  # يُعالج في فرعي الجداول أدناه
        # نصوص وكود وJSON وMD — قراءة مباشرة
        elif mime.startswith("text/") or mime == "application/json" or ext in _TEXT_EXTS:
            txt = p.read_text(encoding="utf-8", errors="replace")
            out["content"] = txt
            out["pages"] = txt.count("\n") + 1          # أسطر
            return out

        # PDF — pdfplumber ثم PyPDF2 كاحتياط
        if mime == "application/pdf" or ext == "pdf":
            try:
                import pdfplumber
                with pdfplumber.open(str(p)) as pdf:
                    out["pages"] = len(pdf.pages)
                    out["content"] = "\n\n".join(
                        (pg.extract_text() or "") for pg in pdf.pages)
                return out
            except ImportError:
                pass
            try:
                from PyPDF2 import PdfReader
                r = PdfReader(str(p))
                out["pages"] = len(r.pages)
                out["content"] = "\n\n".join((pg.extract_text() or "") for pg in r.pages)
                return out
            except ImportError:
                return {"content": "", "pages": 0,
                        "error": "استخراج PDF يتطلب: pip install pdfplumber"}

        # DOCX — python-docx
        if ext == "docx" or mime.endswith("wordprocessingml.document"):
            try:
                import docx
                d = docx.Document(str(p))
                paras = [par.text for par in d.paragraphs]
                out["content"] = "\n".join(paras)
                out["pages"] = len([x for x in paras if x.strip()])
                return out
            except ImportError:
                return {"content": "", "pages": 0,
                        "error": "استخراج DOCX يتطلب: pip install python-docx"}

        # XLSX / CSV — pandas (وcsv القياسية كاحتياط للـCSV)
        if ext in ("xlsx", "xls") or mime.endswith("spreadsheetml.sheet"):
            try:
                import pandas as pd
                sheets = pd.read_excel(str(p), sheet_name=None)
                parts, rows = [], 0
                for sname, df in sheets.items():
                    rows += len(df)
                    parts.append(f"[ورقة: {sname}]\n" + df.to_string(index=False))
                out["content"] = "\n\n".join(parts)
                out["pages"] = rows
                return out
            except ImportError:
                return {"content": "", "pages": 0,
                        "error": "استخراج XLSX يتطلب: pip install pandas openpyxl"}
        if ext == "csv" or mime == "text/csv":
            try:
                import pandas as pd
                df = pd.read_csv(str(p))
                out["content"] = df.to_string(index=False)
                out["pages"] = len(df)
                return out
            except ImportError:
                import csv as _csv
                with open(p, encoding="utf-8", errors="replace", newline="") as f:
                    rows = list(_csv.reader(f))
                out["content"] = "\n".join("\t".join(r) for r in rows)
                out["pages"] = max(len(rows) - 1, 0)
                return out

        # PowerPoint — python-pptx (نصوص الشرائح)
        if ext in ("pptx", "ppt") or mime.endswith("presentationml.presentation"):
            try:
                from pptx import Presentation
                prs = Presentation(str(p))
                slides = []
                for i, slide in enumerate(prs.slides, 1):
                    texts = [sh.text for sh in slide.shapes
                             if getattr(sh, "has_text_frame", False) and sh.text.strip()]
                    slides.append(f"[شريحة {i}]\n" + "\n".join(texts))
                out["content"] = "\n\n".join(slides)
                out["pages"] = len(prs.slides._sldIdLst) if hasattr(prs.slides, "_sldIdLst") \
                    else len(slides)
                return out
            except ImportError:
                return {"content": "", "pages": 0,
                        "error": "استخراج PowerPoint يتطلب: pip install python-pptx"}

        # ZIP — يقرأ ما بداخله: نصوص الملفات النصية + قائمة الصور والملفات الثنائية
        if ext == "zip" or mime in ("application/zip", "application/x-zip-compressed"):
            return _extract_zip(p)

        return {"content": "", "pages": 0, "error": "نوع الملف غير مدعوم"}
    except Exception as e:
        return {"content": "", "pages": 0, "error": f"فشل الاستخراج: {str(e)[:80]}"}


_ZIP_IMAGE_EXTS = {"png", "jpg", "jpeg", "gif", "webp", "bmp", "svg", "tif", "tiff"}
_ZIP_PER_FILE_LIMIT = 8000     # حرف لكل ملف نصي داخل الأرشيف
_ZIP_TOTAL_LIMIT = 60000       # سقف إجمالي المحتوى المستخرج من الأرشيف


def _extract_zip(p: Path) -> Dict[str, Any]:
    """يقرأ محتوى أرشيف ZIP: نصوص الملفات النصية + جرد الصور والملفات الثنائية."""
    import zipfile
    try:
        with zipfile.ZipFile(p) as zf:
            entries = [i for i in zf.infolist() if not i.is_dir()]
            parts: List[str] = [f"[أرشيف ZIP يحتوي {len(entries)} ملف]"]
            images: List[str] = []
            total = 0
            for info in entries:
                name = info.filename
                ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
                size = info.file_size
                if ext in _ZIP_IMAGE_EXTS:
                    images.append(f"  🖼️ {name} ({round(size/1024, 1)}KB)")
                    continue
                if ext in _TEXT_EXTS and total < _ZIP_TOTAL_LIMIT:
                    try:
                        raw = zf.read(info)[: _ZIP_PER_FILE_LIMIT * 2]
                        txt = raw.decode("utf-8", errors="replace")[:_ZIP_PER_FILE_LIMIT]
                        block = f"\n── [{name}] ──\n{txt}"
                        parts.append(block)
                        total += len(block)
                    except Exception:
                        parts.append(f"\n── [{name}] — تعذّرت القراءة ──")
                else:
                    parts.append(f"  📎 {name} ({round(size/1024, 1)}KB)")
            if images:
                parts.append("\n[الصور داخل الأرشيف]\n" + "\n".join(images))
            content = "\n".join(parts)
            if total >= _ZIP_TOTAL_LIMIT:
                content += "\n\n[تم اقتطاع بقية المحتوى — الأرشيف كبير]"
            return {"content": content, "pages": len(entries)}
    except zipfile.BadZipFile:
        return {"content": "", "pages": 0, "error": "أرشيف ZIP تالف أو غير صالح"}
    except RuntimeError as e:  # أرشيف مشفّر بكلمة مرور
        return {"content": "", "pages": 0,
                "error": "أرشيف محمي بكلمة مرور — تعذّر فتحه" if "password" in str(e).lower()
                else f"فشل قراءة الأرشيف: {str(e)[:60]}"}
    except Exception as e:
        return {"content": "", "pages": 0, "error": f"فشل قراءة الأرشيف: {str(e)[:60]}"}


def build_context_block(saved: List[Dict[str, Any]]) -> str:
    """يبني كتلة نصية عن المرفقات تُلحق برسالة المستخدم للنموذج."""
    if not saved:
        return ""
    parts = ["\n\n[مرفقات المستخدم مع هذه الرسالة]"]
    for r in saved:
        size_kb = round((r.get("size") or 0) / 1024, 1)
        head = f"• {r['name']} ({r.get('type') or 'ملف'}, {size_kb}KB)"
        if not r.get("saved"):
            parts.append(head + f" — لم يُحفظ: {r.get('error', '?')}")
            continue
        if r.get("text"):
            parts.append(head + " — المحتوى:\n```\n" + r["text"] + "\n```")
        else:
            parts.append(head + f" — حُفظ في {r['path']} (ملف ثنائي؛ صفه للمستخدم "
                                "واذكر أنه محفوظ لديه)")
    return "\n".join(parts)


def upload_files(data: Dict[str, Any]) -> Tuple[Dict[str, Any], int]:
    """معالج POST /api/files/upload — {files: [{name,type,size,data}]}

    يعيد لكل ملف: path, name, saved_name, mime_type, size, content,
    preview (أول 200 حرف), pages — والصور بلا content (image: true).
    """
    files = (data or {}).get("files") or []
    if not files:
        return {"success": False, "error": "لا ملفات في الطلب"}, 400
    saved = save_attachments(files)
    out = []
    for r in saved:
        content = r.get("content") or ""
        out.append({
            "path": r.get("path"), "name": r.get("name"),
            "saved_name": r.get("saved_name"),
            "mime_type": r.get("type") or "application/octet-stream",
            "size": r.get("size", 0), "saved": r.get("saved", False),
            "content": content, "preview": content[:200],
            "pages": r.get("pages", 0),
            "image": bool(r.get("image")),
            "error": r.get("error") or r.get("extract_error"),
        })
    return {"success": True, "files": out}, 200


# ── قائمة/تنزيل/حذف/استخراج الملفات المحفوظة ─────────────────────
def list_files() -> Tuple[Dict[str, Any], int]:
    """GET /api/files/list — كل ملفات uploads (الأحدث أولاً)."""
    import mimetypes
    UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
    files = []
    for f in sorted(UPLOADS_DIR.iterdir(), key=lambda x: x.stat().st_mtime, reverse=True):
        if f.is_file():
            files.append({
                "name":      f.name,
                "size":      f.stat().st_size,
                "modified":  int(f.stat().st_mtime),
                "mime_type": mimetypes.guess_type(f.name)[0] or "application/octet-stream",
            })
    return {"success": True, "files": files}, 200


def _resolve_upload(filename: str) -> Path:
    """مسار آمن داخل uploads فقط (يمنع الاجتياز)."""
    safe = _safe_name(filename)
    p = (UPLOADS_DIR / safe).resolve()
    if not str(p).startswith(str(UPLOADS_DIR.resolve())):
        raise ValueError("مسار غير مسموح")
    return p


def file_bytes(filename: str):
    """GET /api/files/download/<filename> — يعيد (bytes, mime, name) أو None."""
    import mimetypes
    try:
        p = _resolve_upload(filename)
    except ValueError:
        return None
    if not p.exists() or not p.is_file():
        return None
    mime = mimetypes.guess_type(p.name)[0] or "application/octet-stream"
    return p.read_bytes(), mime, p.name


def extract_for(filename: str) -> Tuple[Dict[str, Any], int]:
    """GET /api/files/extract/<filename> — محتوى ملف محفوظ للعارض."""
    import mimetypes
    try:
        p = _resolve_upload(filename)
    except ValueError:
        return {"success": False, "error": "مسار غير مسموح"}, 400
    if not p.exists():
        return {"success": False, "error": "الملف غير موجود"}, 404
    mime = mimetypes.guess_type(p.name)[0] or "application/octet-stream"
    res = extract_content(p, mime)
    return {"success": True, "name": p.name, "mime_type": mime,
            "size": p.stat().st_size, "content": res.get("content", ""),
            "pages": res.get("pages", 0), "image": bool(res.get("image")),
            "error": res.get("error")}, 200


def delete_file(filename: str) -> Tuple[Dict[str, Any], int]:
    """DELETE /api/files/delete/<filename>"""
    try:
        p = _resolve_upload(filename)
    except ValueError:
        return {"success": False, "error": "مسار غير مسموح"}, 400
    if p.exists() and p.is_file():
        p.unlink()
        logger.info("Upload deleted: %s", p.name)
        return {"success": True}, 200
    return {"success": False, "error": "الملف غير موجود"}, 404
